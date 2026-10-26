#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
inspecionar_insumos_governanca.py — Auditor Determinístico de Insumos de Governança (PRODEB / MongoDB / AWS SageMaker)

Objetivo:
  Inspecionar planilhas de insumos (.xlsx com abas 'Dados - Insumos' e 'Dicionário de Dados')
  e amostras de dados (.csv) enviadas pelos órgãos clientes na pasta `entrega/governanca_dados/`
  (ou em um diretório/arquivo informado via linha de comando).

Regra de Ouro (Gate 0.5 — Responsabilidade do Cliente na Origem):
  A equipe técnica atua exclusivamente no transporte para o MongoDB e publicação no AWS Glue /
  Amazon SageMaker Catalog (ex-DataZone). Se o cliente enviar planilhas incompletas, sem classificação
  LGPD, com textos evasivos ("ver com sap", "xxxx") ou arquivos de dados com registros vazios /
  colunas fantasmas (;;;), este script REPROVA o insumo e gera o Relatório Formal de Devolução
  ao Cliente para saneamento na origem.

Uso:
  python core/scripts_comuns/inspecionar_insumos_governanca.py [caminho_opcional] [--salvar-relatorio caminho.md]
"""

import os
import re
import sys
import csv
import unicodedata
import argparse
from pathlib import Path
from datetime import datetime

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

try:
    from gerador_artefatos_visuais import (
        exportar_markdown_para_docx,
        exportar_markdown_para_pdf,
        extrair_texto_pdf_nativo,
        ler_planilha_ods,
        ler_documento_docx,
        ler_documento_odt,
        escanear_pii_em_texto_ou_linhas,
    )
except Exception:
    exportar_markdown_para_docx = None
    exportar_markdown_para_pdf = None
    extrair_texto_pdf_nativo = None
    ler_planilha_ods = None
    ler_documento_docx = None
    ler_documento_odt = None
    escanear_pii_em_texto_ou_linhas = None

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = Path(__file__).resolve().parent.parent.parent
DEFAULT_ENTREGA_DIR = BASE_DIR / "entrega" / "governanca_dados"

VALORES_EVASIVOS = {
    "", "-", "--", "---", "?", "??", "n/a", "na", "null", "none", "tbd",
    "a definir", "pendente", "ver com sap", "verificar", "nao sei", "não sei",
    "xxxx", "xxx", "xx", "teste", "todo", "preencher"
}

PADROES_EVASIVOS_REGEX = [
    r"\bver com\b",
    r"\bver sap\b",
    r"^versao banco de dados$",
    r"\ba definir\b",
    r"\bverificar\b",
    r"^x{2,}$",
    r"^\?+$",
    r"\bpendente\b",
    r"\bnao informado\b",
]

KEYWORDS_PII_CAMPO = [
    "cpf", "cnpj", "rg", "identidade", "cnh", "renavam", "placa", "chassi",
    "titulo_eleitor", "pis", "pasep", "nis", "carteira_trabalho", "ctps",
    "nome", "nom_", "vorna", "nachn", "mae", "pai", "filiacao", "conjuge",
    "email", "e_mail", "usrid", "telefone", "celular", "tel", "ddd",
    "endereco", "logradouro", "stras", "cep", "bairro", "ort02",
    "matricula", "pernr", "mat_principal", "no_pessoal",
    "banco", "agencia", "conta_corrente", "pix",
    "salario", "sal_principal", "soma_sal", "remuneracao", "vencimento", "provento", "pensao"
]

KEYWORDS_SENSIVEL_ART11 = [
    "cid", "doenca", "patologia", "diagnostico", "deficiencia", "pcd",
    "invalidez", "laudo", "junta_medica", "prontuario", "saude",
    "biometria", "raca", "etnia", "religiao", "sindicato", "sindical",
    "filiacao_partidaria", "orientacao_sexual", "obito", "causa_mortis"
]

PREFIXOS_PRODEB_VALIDOS = (
    "cod_", "nom_", "des_", "num_", "ind_", "dtc_", "dth_", "val_",
    "qtd_", "prc_", "sts_", "tip_", "ref_", "lst_", "jsn_", "sig_",
    "txt_", "hor_"
)


def remover_acentos(texto: str) -> str:
    nfkd = unicodedata.normalize("NFKD", str(texto))
    return "".join(c for c in nfkd if not unicodedata.combining(c))


def eh_valor_vazio_ou_evasivo(val) -> tuple[bool, str]:
    if val is None:
        return True, "VAZIO (em branco)"
    limpo = str(val).strip()
    if not limpo:
        return True, "VAZIO (em branco)"
    baixo = remover_acentos(limpo).lower()
    if baixo in VALORES_EVASIVOS:
        return True, f"VALOR EVASIVO/INVÁLIDO ('{limpo}')"
    for pat in PADROES_EVASIVOS_REGEX:
        if re.search(pat, baixo):
            return True, f"TEXTO EVASIVO/INCOMPLETO ('{limpo}')"
    return False, ""


def inferir_prefixo_e_nome_mongo(nome_campo: str, descricao: str, tipo_origem: str) -> tuple[str, str]:
    """
    Converte um nome de coluna legado para o padrão PRODEB MongoDB:
      - snake_case minúsculo sem acentos
      - prefixo semântico de 3 letras (cod_, nom_, val_, dtc_, etc.)
      - tipo BSON rigoroso (string, int, long, decimal, bool, date)
    """
    bruto = remover_acentos(nome_campo).strip().lower()
    bruto = re.sub(r"[^a-z0-9]+", "_", bruto).strip("_")
    desc_lower = remover_acentos(descricao or "").strip().lower()
    tipo_lower = remover_acentos(tipo_origem or "").strip().lower()

    # 1. Determinar tipo BSON e prefixo semântico
    eh_doc_codigo = any(
        k in bruto
        for k in ["cpf", "cnpj", "cep", "rg", "cnh", "renavam", "placa", "pis", "pasep", "nis", "matricula", "pernr", "no_pessoal", "telefone", "celular", "tel", "codigo", "cod_", "_id", "id_", "protocolo", "processo"]
    )
    eh_monetario = (
        any(k in bruto for k in ["val_", "valor", "vlr", "sal_", "salario", "remuneracao", "vencimento", "provento", "desconto", "liquido", "bruto", "custo", "preco", "tarifa", "montante", "soma_sal"])
        and "data" not in tipo_lower and "date" not in tipo_lower
    )
    eh_percentual = any(
        k in bruto for k in ["prc_", "perc", "percentual", "aliquota", "indice", "taxa_perc"]
    )
    eh_datahora = any(
        k in bruto or "timestamp" in tipo_lower or "datetime" in tipo_lower
        for k in ["dth_", "data_hora", "dt_hr", "timestamp", "created_at", "updated_at", "dh_"]
    )
    eh_data = any(
        k in bruto or "date" in tipo_lower or tipo_lower == "data"
        for k in ["dtc_", "dt_", "data_", "_data", "dat_", "nascimento", "admissao", "vigencia", "competencia", "emissao", "validade", "obito", "begda", "gbdat", "mes_ref"]
    )
    eh_hora = any(k in bruto for k in ["hor_", "hora_", "hr_"])
    eh_booleano = any(
        k in bruto or "bool" in tipo_lower or "bit" in tipo_lower
        for k in ["ind_", "flg_", "flag_", "ativo", "inativo", "possui_", "eh_", "is_"]
    )
    eh_status = any(k in bruto for k in ["sts_", "status", "situacao", "sit_"])
    eh_sigla = any(k in bruto for k in ["sig_", "sg_", "sigla", "uf"])
    eh_qtd = any(k in bruto for k in ["qtd_", "qtde", "quant"])
    eh_nome = any(k in bruto for k in ["nom_", "nm_", "nome", "vorna", "nachn"])
    eh_desc = any(k in bruto for k in ["des_", "ds_", "desc_", "descricao", "motivo", "observacao", "mntxt", "mgtxt"])
    eh_tipo = any(k in bruto for k in ["tip_", "tp_", "tipo_"])
    eh_num = any(k in bruto for k in ["num_", "nr_", "nro_", "numero", "ano", "mes", "idade"])

    # Escolha do prefixo
    if bruto.startswith(PREFIXOS_PRODEB_VALIDOS) and not (bruto.startswith("val_") and eh_data):
        nome_final = bruto
        prefixo = bruto[:4]
    else:
        sem_pref_legado = re.sub(r"^(cd|id|nm|ds|nr|vl|dt|dh|tp|sg|fl|st|qt|sal)_", "", bruto)
        if eh_datahora:
            prefixo = "dth_"
        elif eh_data:
            prefixo = "dtc_"
        elif eh_monetario:
            prefixo = "val_"
        elif eh_percentual:
            prefixo = "prc_"
        elif eh_hora:
            prefixo = "hor_"
        elif eh_booleano:
            prefixo = "ind_"
        elif eh_status:
            prefixo = "sts_"
        elif eh_sigla:
            prefixo = "sig_"
        elif eh_qtd:
            prefixo = "qtd_"
        elif eh_nome:
            prefixo = "nom_"
        elif eh_desc:
            prefixo = "des_"
        elif eh_tipo:
            prefixo = "tip_"
        elif eh_doc_codigo:
            prefixo = "cod_"
        elif eh_num:
            prefixo = "num_"
        else:
            prefixo = "des_"
        nome_final = f"{prefixo}{sem_pref_legado}"

    # Escolha do tipo BSON segundo padrão PRODEB
    if prefixo in ("dtc_", "dth_") or eh_data or eh_datahora:
        bson_type = "date"
    elif prefixo in ("val_", "prc_") or eh_monetario or eh_percentual:
        bson_type = "decimal"
    elif prefixo == "ind_" or eh_booleano:
        bson_type = "bool"
    elif prefixo == "cod_" or eh_doc_codigo:
        bson_type = "string"
    elif prefixo in ("qtd_", "num_") or any(t in tipo_lower for t in ["int", "integer", "number", "numeric", "numérico", "numerico", "bigint", "smallint"]):
        if any(k in bruto for k in ["cpf", "cnpj", "cep", "rg", "telefone", "celular", "matricula", "processo", "beneficiario"]):
            bson_type = "string"
        elif "texto" in tipo_lower or "char" in tipo_lower or "string" in tipo_lower:
            bson_type = "string"
        elif "dec" in tipo_lower or "float" in tipo_lower or "double" in tipo_lower or "real" in tipo_lower or "money" in tipo_lower:
            bson_type = "decimal"
        else:
            bson_type = "int"
    elif any(t in tipo_lower for t in ["dec", "float", "double", "real", "money", "currency"]):
        bson_type = "decimal"
    else:
        bson_type = "string"

    return nome_final, bson_type


def detectar_risco_lgpd(nome_campo: str, descricao: str) -> tuple[bool, bool, str]:
    campo_norm = re.sub(r"[^a-z0-9]+", "_", remover_acentos(nome_campo).lower()).strip("_")
    desc_norm = re.sub(r"[^a-z0-9]+", "_", remover_acentos(descricao or "").lower()).strip("_")

    for kw in KEYWORDS_SENSIVEL_ART11:
        if re.search(rf"(^|_){re.escape(kw)}($|_)", campo_norm) or re.search(rf"(^|_){re.escape(kw)}($|_)", desc_norm):
            return True, True, f"Dado Pessoal Sensível (Art. 11 LGPD — '{kw}')"
    for kw in KEYWORDS_PII_CAMPO:
        if re.search(rf"(^|_){re.escape(kw)}($|_)", campo_norm):
            return True, False, f"Dado Pessoal / PII (Art. 5º, I LGPD — '{kw}')"
    for kw in ["cpf", "cnpj", "rg", "nome_completo", "nome_da_mae", "endereco_residencial", "telefone", "e_mail", "salario", "remuneracao"]:
        if re.search(rf"(^|_){re.escape(kw)}($|_)", desc_norm):
            return True, False, f"Dado Pessoal / PII (Art. 5º, I LGPD — '{kw}')"
    return False, False, "Não Pessoal / Operacional"


def inspecionar_xlsx(caminho_xlsx: Path) -> dict:
    resultado = {
        "arquivo": caminho_xlsx.name,
        "caminho": str(caminho_xlsx),
        "tipo_arquivo": "XLSX",
        "bloqueios_devolucao": [],
        "alertas_lgpd": [],
        "metadados_insumos": {},
        "de_para_campos": [],
        "resumo": {}
    }

    if not HAS_OPENPYXL:
        resultado["bloqueios_devolucao"].append("Biblioteca 'openpyxl' não disponível no ambiente Python.")
        return resultado

    try:
        wb = openpyxl.load_workbook(caminho_xlsx, data_only=True)
    except Exception as exc:
        resultado["bloqueios_devolucao"].append(f"Arquivo Excel corrompido ou ilegível: {exc}")
        return resultado

    abas = wb.sheetnames
    aba_insumos = next((s for s in abas if "insumo" in remover_acentos(s).lower()), None)
    aba_dic = next((s for s in abas if "dicionario" in remover_acentos(s).lower()), None)

    if not aba_insumos:
        resultado["bloqueios_devolucao"].append(
            f"Aba obrigatória 'Dados - Insumos' não encontrada (abas existentes: {abas})."
        )
    if not aba_dic:
        resultado["bloqueios_devolucao"].append(
            f"Aba obrigatória 'Dicionário de Dados' não encontrada (abas existentes: {abas})."
        )

    # 1. Auditar Aba 'Dados - Insumos' (suporta layout 2 colunas [Objeto|Descrição] ou 3 colunas [Categoria|Objeto|Descrição])
    if aba_insumos:
        ws_ins = wb[aba_insumos]
        linhas_insumos = 0
        col_objeto = 1
        col_valor = 2

        # Detectar se coluna 2 é 'Objeto' e coluna 3 é 'Descrição'
        for r_probe in range(1, min(6, ws_ins.max_row + 1)):
            c1 = remover_acentos(ws_ins.cell(row=r_probe, column=1).value or "").strip().lower()
            c2 = remover_acentos(ws_ins.cell(row=r_probe, column=2).value or "").strip().lower()
            if "categoria" in c1 or "objeto" in c2:
                col_objeto = 2
                col_valor = 3
                break

        for row_idx in range(1, ws_ins.max_row + 1):
            chave = ws_ins.cell(row=row_idx, column=col_objeto).value
            valor = ws_ins.cell(row=row_idx, column=col_valor).value
            if chave is None and valor is None:
                continue
            chave_str = str(chave or "").strip()
            if not chave_str or chave_str.lower() in ("objeto", "item", "campo", "atributo") or chave_str.lower().startswith("dados - insumos"):
                continue
            linhas_insumos += 1
            invalido, motivo = eh_valor_vazio_ou_evasivo(valor)
            resultado["metadados_insumos"][chave_str] = str(valor).strip() if valor is not None else ""
            if invalido:
                resultado["bloqueios_devolucao"].append(
                    f"[Aba '{aba_insumos.strip()}' | Linha {row_idx}] Metadado obrigatório '{chave_str}' está {motivo}."
                )
        if linhas_insumos == 0:
            resultado["bloqueios_devolucao"].append(
                f"[Aba '{aba_insumos.strip()}'] Nenhum metadado institucional preenchido."
            )

    # 2. Auditar Aba 'Dicionário de Dados'
    if aba_dic:
        ws_dic = wb[aba_dic]
        header_row = None
        col_map = {}

        for r_idx in range(1, min(15, ws_dic.max_row + 1)):
            vals = [ws_dic.cell(row=r_idx, column=c).value for c in range(1, min(15, ws_dic.max_column + 1))]
            norm_vals = [remover_acentos(v or "").strip().lower() for v in vals]
            if any("campo" in v or "atributo" in v or "coluna" in v for v in norm_vals):
                header_row = r_idx
                for c_idx, nv in enumerate(norm_vals, start=1):
                    if ("campo" in nv or "coluna" in nv or "atributo" in nv) and "nome" not in col_map:
                        col_map["nome"] = c_idx
                    elif "descri" in nv and "descricao" not in col_map:
                        col_map["descricao"] = c_idx
                    elif "tipo" in nv and "tipo" not in col_map:
                        col_map["tipo"] = c_idx
                    elif "anonim" in nv and "anonimizar" not in col_map:
                        col_map["anonimizar"] = c_idx
                    elif "sensibilidade" in nv and "sensibilidade" not in col_map:
                        col_map["sensibilidade"] = c_idx
                    elif "sensivel" in nv and "sensivel" not in col_map:
                        col_map["sensivel"] = c_idx
                break

        if not header_row or "nome" not in col_map:
            resultado["bloqueios_devolucao"].append(
                f"[Aba '{aba_dic}'] Cabeçalho padrão ('Nome do Campo', 'Descrição', 'Tipo', 'Anonimizar?', 'Sensibilidade', 'Sensível?') não localizado."
            )
        else:
            total_campos = 0
            campos_sem_desc = 0
            campos_sem_tipo = 0
            campos_sem_lgpd = 0

            for r_idx in range(header_row + 1, ws_dic.max_row + 1):
                nome_val = ws_dic.cell(row=r_idx, column=col_map["nome"]).value
                desc_val = ws_dic.cell(row=r_idx, column=col_map.get("descricao", 2)).value if "descricao" in col_map else None
                tipo_val = ws_dic.cell(row=r_idx, column=col_map.get("tipo", 3)).value if "tipo" in col_map else None
                anon_val = ws_dic.cell(row=r_idx, column=col_map.get("anonimizar", 4)).value if "anonimizar" in col_map else None
                sensib_val = ws_dic.cell(row=r_idx, column=col_map.get("sensibilidade", 5)).value if "sensibilidade" in col_map else None
                sensiv_val = ws_dic.cell(row=r_idx, column=col_map.get("sensivel", 6)).value if "sensivel" in col_map else None

                if all(v is None or str(v).strip() == "" for v in [nome_val, desc_val, tipo_val, anon_val, sensib_val, sensiv_val]):
                    continue

                total_campos += 1
                nome_str = str(nome_val or "").strip()

                inv_nome, mot_nome = eh_valor_vazio_ou_evasivo(nome_val)
                inv_desc, mot_desc = eh_valor_vazio_ou_evasivo(desc_val)
                inv_tipo, mot_tipo = eh_valor_vazio_ou_evasivo(tipo_val)
                inv_anon, _ = eh_valor_vazio_ou_evasivo(anon_val)
                inv_sensib, _ = eh_valor_vazio_ou_evasivo(sensib_val)
                inv_sensiv, _ = eh_valor_vazio_ou_evasivo(sensiv_val)

                if inv_nome:
                    resultado["bloqueios_devolucao"].append(
                        f"[Aba '{aba_dic}' | Linha {r_idx}] Registro com 'Nome do Campo' {mot_nome}, mas demais colunas preenchidas."
                    )
                    continue

                if inv_desc:
                    campos_sem_desc += 1
                    resultado["bloqueios_devolucao"].append(
                        f"[Aba '{aba_dic}' | Linha {r_idx} | Campo '{nome_str}'] 'Descrição' está {mot_desc}."
                    )
                if inv_tipo:
                    campos_sem_tipo += 1
                    resultado["bloqueios_devolucao"].append(
                        f"[Aba '{aba_dic}' | Linha {r_idx} | Campo '{nome_str}'] 'Tipo' está {mot_tipo}."
                    )
                if inv_anon or inv_sensib or inv_sensiv:
                    campos_sem_lgpd += 1
                    faltantes = []
                    if inv_anon:
                        faltantes.append("Anonimizar?")
                    if inv_sensib:
                        faltantes.append("Sensibilidade")
                    if inv_sensiv:
                        faltantes.append("Sensível?")
                    resultado["bloqueios_devolucao"].append(
                        f"[Aba '{aba_dic}' | Linha {r_idx} | Campo '{nome_str}'] Classificação LGPD incompleta — coluna(s) vazia(s): {', '.join(faltantes)}."
                    )

                eh_pii, eh_art11, diag_lgpd = detectar_risco_lgpd(nome_str, str(desc_val or ""))
                sensiv_norm = remover_acentos(str(sensiv_val or "")).strip().lower()
                anon_norm = remover_acentos(str(anon_val or "")).strip().lower()

                if (eh_pii or eh_art11) and sensiv_norm in ("nao", "n", "false", "0") and anon_norm in ("nao", "n", "false", "0"):
                    resultado["alertas_lgpd"].append(
                        f"[Linha {r_idx} | Campo '{nome_str}'] Cliente marcou Sensível='Não' e Anonimizar='Não', mas o campo aparenta ser {diag_lgpd}."
                    )

                nom_mongo, bson_mongo = inferir_prefixo_e_nome_mongo(
                    nome_str, str(desc_val or ""), str(tipo_val or "")
                )

                resultado["de_para_campos"].append({
                    "linha": r_idx,
                    "campo_origem": nome_str,
                    "tipo_origem": str(tipo_val or "").strip(),
                    "descricao": str(desc_val or "").strip(),
                    "campo_mongo": nom_mongo,
                    "tipo_bson": bson_mongo,
                    "anonimizar_cliente": str(anon_val or "").strip(),
                    "sensibilidade_cliente": str(sensib_val or "").strip(),
                    "sensivel_cliente": str(sensiv_val or "").strip(),
                    "diagnostico_ia_lgpd": diag_lgpd
                })

            if total_campos == 0:
                resultado["bloqueios_devolucao"].append(
                    f"[Aba '{aba_dic}'] Nenhum campo cadastrado no Dicionário de Dados."
                )

            resultado["resumo"] = {
                "total_campos": total_campos,
                "campos_sem_descricao": campos_sem_desc,
                "campos_sem_tipo": campos_sem_tipo,
                "campos_sem_lgpd": campos_sem_lgpd,
            }

    return resultado


def inspecionar_csv(caminho_csv: Path) -> dict:
    resultado = {
        "arquivo": caminho_csv.name,
        "caminho": str(caminho_csv),
        "tipo_arquivo": "CSV",
        "bloqueios_devolucao": [],
        "alertas_lgpd": [],
        "resumo": {}
    }

    conteudo = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            conteudo = caminho_csv.read_text(encoding=enc)
            break
        except Exception:
            continue

    if conteudo is None:
        resultado["bloqueios_devolucao"].append("Não foi possível decodificar o arquivo CSV (codificação inválida).")
        return resultado

    linhas = [ln for ln in conteudo.splitlines() if ln.strip() != ""]
    if not linhas:
        resultado["bloqueios_devolucao"].append("Arquivo CSV está completamente vazio (0 linhas).")
        return resultado

    primeira_linha = linhas[0]
    delimitador = ";" if primeira_linha.count(";") >= primeira_linha.count(",") else ","

    leitor = list(csv.reader(linhas, delimiter=delimitador))
    cabecalho = leitor[0]
    dados = leitor[1:]

    # 1. Verificar colunas fantasmas / vazias no cabeçalho (ex: ;;;;; no final)
    cols_vazias_idx = [i + 1 for i, col in enumerate(cabecalho) if str(col).strip() == ""]
    cols_validas_idx = [i for i, col in enumerate(cabecalho) if str(col).strip() != ""]

    if cols_vazias_idx:
        resultado["bloqueios_devolucao"].append(
            f"[Cabeçalho CSV] Encontradas {len(cols_vazias_idx)} coluna(s) fantasma(s)/sem nome no cabeçalho (posições: {cols_vazias_idx[:10]}). O cliente deve remover delimitadores extras (';;;') na exportação."
        )

    if not dados:
        resultado["bloqueios_devolucao"].append(
            "[Amostra CSV] O arquivo possui apenas cabeçalho e nenhum registro de dados (0 linhas de dados)."
        )
        return resultado

    # 2. Auditar linhas vazias, desalinhadas ou com chaves/campos obrigatórios vazios
    linhas_totalmente_vazias = 0
    linhas_chave_vazia = 0
    linhas_desalinhadas = 0
    vazios_por_coluna = {cabecalho[i].strip(): 0 for i in cols_validas_idx}

    for row_num, row in enumerate(dados, start=2):
        if len(row) != len(cabecalho):
            linhas_desalinhadas += 1

        valores_uteis = [str(row[i]).strip() if i < len(row) else "" for i in cols_validas_idx]
        if all(v == "" or v.upper() == "NULL" for v in valores_uteis):
            linhas_totalmente_vazias += 1
            continue

        # Primeira coluna válida (geralmente ID / Matrícula / CPF / Nome / Código) vazia
        if valores_uteis and (valores_uteis[0] == "" or valores_uteis[0].upper() == "NULL"):
            linhas_chave_vazia += 1

        for idx_col in cols_validas_idx:
            nome_col = cabecalho[idx_col].strip()
            val_cel = str(row[idx_col]).strip() if idx_col < len(row) else ""
            if val_cel == "" or val_cel.upper() == "NULL":
                vazios_por_coluna[nome_col] += 1

    if linhas_totalmente_vazias > 0:
        resultado["bloqueios_devolucao"].append(
            f"[Amostra CSV] Encontrados {linhas_totalmente_vazias} registro(s) com todos os campos em branco/NULL."
        )
    if linhas_chave_vazia > 0:
        nome_chave = cabecalho[cols_validas_idx[0]].strip() if cols_validas_idx else "Coluna 1"
        resultado["bloqueios_devolucao"].append(
            f"[Amostra CSV] Encontrados {linhas_chave_vazia} registro(s) com a coluna identificadora inicial ('{nome_chave}') vazia."
        )
    if linhas_desalinhadas > 0:
        resultado["bloqueios_devolucao"].append(
            f"[Amostra CSV] Encontradas {linhas_desalinhadas} linha(s) com número de colunas diferente do cabeçalho (quebra de delimitador '{delimitador}')."
        )

    total_linhas_dados = len(dados)
    colunas_100_vazias = [
        col for col, qtd_vazio in vazios_por_coluna.items() if qtd_vazio == total_linhas_dados
    ]
    if colunas_100_vazias:
        resultado["bloqueios_devolucao"].append(
            f"[Amostra CSV] Coluna(s) com 100% dos registros vazios na amostra ({len(colunas_100_vazias)}): {', '.join(colunas_100_vazias[:10])}."
        )

    # 3. Scanner de PII Escondido no Conteúdo das Linhas (Redact / LGPD Scanner)
    if escanear_pii_em_texto_ou_linhas is not None:
        achados_pii = escanear_pii_em_texto_ou_linhas(leitor, nome_contexto=caminho_csv.name)
        for ach in achados_pii:
            resultado["alertas_lgpd"].append(
                f"[Scanner Conteúdo CSV | Linha {ach['linha']} | Coluna '{ach['coluna']}'] "
                f"Detectado {ach['tipo_pii']} (`{ach['valor_mascarado']}`) — Classificação recomendada: "
                f"{ach['classificacao_sugerida']} ({ach['base_legal_sugerida']})."
            )

    resultado["resumo"] = {
        "delimitador": delimitador,
        "total_colunas_cabecalho": len(cabecalho),
        "colunas_validas": len(cols_validas_idx),
        "colunas_fantasmas_vazias": len(cols_vazias_idx),
        "total_registros_amostra": total_linhas_dados,
        "colunas_100_porcento_vazias": colunas_100_vazias,
    }
    return resultado


def inspecionar_ods(caminho_ods: Path) -> dict:
    """Inspeciona planilhas LibreOffice Calc (.ods) aplicando as regras do Gate 0.5 PRODEB."""
    resultado = {
        "arquivo": caminho_ods.name,
        "caminho": str(caminho_ods),
        "tipo_arquivo": "ODS",
        "bloqueios_devolucao": [],
        "alertas_lgpd": [],
        "metadados_insumos": {},
        "de_para_campos": [],
        "resumo": {},
    }
    if ler_planilha_ods is None:
        resultado["bloqueios_devolucao"].append("Leitor nativo .ods indisponível.")
        return resultado
    try:
        rows = ler_planilha_ods(str(caminho_ods))
    except Exception as exc:
        resultado["bloqueios_devolucao"].append(f"Planilha ODS corrompido ou ilegível: {exc}")
        return resultado

    if not rows:
        resultado["bloqueios_devolucao"].append("Arquivo .ods corrompido, vazio ou sem abas legíveis.")
        return resultado

    for r_idx, row in enumerate(rows[1:], start=2):
        if not row or all(not str(c or "").strip() for c in row):
            continue
        campo = str(row[0] if len(row) > 0 else "").strip()
        desc = str(row[1] if len(row) > 1 else "").strip()
        tipo = str(row[2] if len(row) > 2 else "").strip()
        if not campo or campo.lower() in ("campo", "coluna", "atributo", "nome_campo"):
            continue
        inv_d, mot_d = eh_valor_vazio_ou_evasivo(desc)
        if inv_d:
            resultado["bloqueios_devolucao"].append(
                f"[Planilha ODS | Linha {r_idx}] Campo '{campo}' com Descrição {mot_d}."
            )
        campo_mongo, tipo_bson = inferir_prefixo_e_nome_mongo(campo, desc, tipo)
        _, _, diag_lgpd = detectar_risco_lgpd(campo, desc)
        resultado["de_para_campos"].append(
            {
                "linha": r_idx,
                "campo_origem": campo,
                "tipo_origem": tipo,
                "campo_mongo": campo_mongo,
                "tipo_bson": tipo_bson,
                "diagnostico_ia_lgpd": diag_lgpd,
            }
        )

    if escanear_pii_em_texto_ou_linhas is not None:
        for ach in escanear_pii_em_texto_ou_linhas(rows, nome_contexto=caminho_ods.name):
            resultado["alertas_lgpd"].append(
                f"[Scanner Conteúdo ODS | Linha {ach['linha']} | Coluna '{ach['coluna']}'] "
                f"Detectado {ach['tipo_pii']} (`{ach['valor_mascarado']}`) — {ach['classificacao_sugerida']}."
            )
    return resultado


def inspecionar_pdf(caminho_pdf: Path) -> dict:
    """Inspeciona documentos/ofícios PDF de governança de forma resiliente (`panic_hunt`)."""
    resultado = {
        "arquivo": caminho_pdf.name,
        "caminho": str(caminho_pdf),
        "tipo_arquivo": "PDF",
        "bloqueios_devolucao": [],
        "alertas_lgpd": [],
        "metadados_insumos": {},
        "de_para_campos": [],
        "resumo": {},
    }
    if extrair_texto_pdf_nativo is None:
        return resultado
    try:
        dados_raw = caminho_pdf.read_bytes()
        if not dados_raw.startswith(b"%PDF"):
            resultado["bloqueios_devolucao"].append("Documento PDF inválido ou truncado (Cabeçalho %PDF ausente).")
            return resultado
        texto = (extrair_texto_pdf_nativo(str(caminho_pdf)) or "").strip()
    except Exception as exc:
        resultado["bloqueios_devolucao"].append(f"Documento PDF corrompido ou ilegível: {exc}")
        return resultado

    resultado["resumo"] = {"caracteres_extraidos": len(texto)}
    if not texto:
        resultado["bloqueios_devolucao"].append(
            "Documento PDF sem camada de texto extraível (imagem digitalizada sem OCR ou stream vazio)."
        )
    elif escanear_pii_em_texto_ou_linhas is not None:
        for ach in escanear_pii_em_texto_ou_linhas(texto, nome_contexto=caminho_pdf.name):
            resultado["alertas_lgpd"].append(
                f"[Scanner Conteúdo PDF | Linha {ach['linha']}] Detectado {ach['tipo_pii']} (`{ach['valor_mascarado']}`) — {ach['classificacao_sugerida']}."
            )
    return resultado


def inspecionar_doc_texto(caminho_doc: Path) -> dict:
    """Inspeciona documentos Word (.docx) e LibreOffice Writer (.odt) na governança."""
    suf = caminho_doc.suffix.lower()
    fmt = "DOCX" if suf == ".docx" else "ODT"
    resultado = {
        "arquivo": caminho_doc.name,
        "caminho": str(caminho_doc),
        "tipo_arquivo": fmt,
        "bloqueios_devolucao": [],
        "alertas_lgpd": [],
        "metadados_insumos": {},
        "de_para_campos": [],
        "resumo": {},
    }
    leitor_fn = ler_documento_docx if suf == ".docx" else ler_documento_odt
    if leitor_fn is None:
        return resultado
    try:
        info = leitor_fn(str(caminho_doc))
    except Exception as exc:
        resultado["bloqueios_devolucao"].append(f"Documento {fmt} corrompido ou ilegível: {exc}")
        return resultado

    texto = (info.get("texto") or "").strip()
    tabelas = info.get("tabelas") or []
    resultado["resumo"] = {"caracteres_extraidos": len(texto), "tabelas": len(tabelas)}
    if not texto and not tabelas:
        resultado["bloqueios_devolucao"].append(f"Documento {fmt} vazio ou sem conteúdo legível.")
        return resultado

    if escanear_pii_em_texto_ou_linhas is not None:
        for ach in escanear_pii_em_texto_ou_linhas(texto, nome_contexto=caminho_doc.name):
            resultado["alertas_lgpd"].append(
                f"[Scanner Conteúdo {fmt} | Linha {ach['linha']}] Detectado {ach['tipo_pii']} (`{ach['valor_mascarado']}`) — {ach['classificacao_sugerida']}."
            )
    return resultado


def gerar_markdown_relatorio(resultados: list[dict]) -> str:
    linhas_md = [
        "# 📋 Relatório de Auditoria de Insumos de Governança & Devolução ao Cliente",
        "",
        f"- **Data/Hora da Inspeção:** `{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}`",
        "- **Política Aplicada:** Gate 0.5 — *Responsabilidade de Qualidade na Origem é do Órgão Cliente (Escopo Técnico: Transporte para MongoDB Enterprise e Publicação no AWS Glue / Amazon SageMaker Catalog)*.",
        ""
    ]

    total_bloqueios = sum(len(r["bloqueios_devolucao"]) for r in resultados)
    status_geral = (
        "❌ REPROVADO NO GATE DE ENTRADA — DEVOLVER AO CLIENTE PARA CORREÇÃO NA ORIGEM"
        if total_bloqueios > 0
        else "✅ APROVADO NO GATE DE ENTRADA — APTO PARA TRANSPORTE (MONGODB ➔ AWS SAGEMAKER CATALOG)"
    )
    linhas_md.append(f"## Veredito Geral: **{status_geral}**")
    linhas_md.append("")

    for r in resultados:
        linhas_md.append(f"---")
        linhas_md.append(f"## 📄 Arquivo: `{r['arquivo']}` ({r['tipo_arquivo']})")
        linhas_md.append("")

        if r["bloqueios_devolucao"]:
            linhas_md.append(f"### 🛑 Pendências Impeditivas para Devolução ao Cliente ({len(r['bloqueios_devolucao'])})")
            linhas_md.append("> O cliente deve sanear os itens abaixo antes do transporte para o MongoDB e AWS SageMaker Catalog:")
            linhas_md.append("")
            for b in r["bloqueios_devolucao"][:40]:
                linhas_md.append(f"- ❌ {b}")
            if len(r["bloqueios_devolucao"]) > 40:
                linhas_md.append(f"- *(... e mais {len(r['bloqueios_devolucao']) - 40} pendências similares no arquivo)*")
            linhas_md.append("")
        else:
            linhas_md.append("### ✅ Sem Pendências Estruturais Bloqueantes")
            linhas_md.append("")

        if r.get("alertas_lgpd"):
            linhas_md.append(f"### ⚠️ Radar Consultivo LGPD / LAI ({len(r['alertas_lgpd'])} alerta(s))")
            for al in r["alertas_lgpd"][:25]:
                linhas_md.append(f"- 🔒 {al}")
            if len(r["alertas_lgpd"]) > 25:
                linhas_md.append(f"- *(... e mais {len(r['alertas_lgpd']) - 25} alertas de LGPD)*")
            linhas_md.append("")

        if r.get("de_para_campos"):
            linhas_md.append("### 🔄 Prévia da Matriz DE/PARA Padrão PRODEB MongoDB (`snake_case` + BSON)")
            linhas_md.append("")
            linhas_md.append("| Linha | Campo Origem | Tipo Origem | Campo Canônico MongoDB | Tipo BSON | Diagnóstico LGPD (IA) |")
            linhas_md.append("| :---: | :--- | :--- | :--- | :---: | :--- |")
            for item in r["de_para_campos"][:30]:
                linhas_md.append(
                    f"| {item['linha']} | `{item['campo_origem']}` | `{item['tipo_origem'] or 'VAZIO'}` | "
                    f"`{item['campo_mongo']}` | `{item['tipo_bson']}` | {item['diagnostico_ia_lgpd']} |"
                )
            if len(r["de_para_campos"]) > 30:
                linhas_md.append(f"| ... | *(+{len(r['de_para_campos']) - 30} campos)* | ... | ... | ... | ... |")
            linhas_md.append("")

    return "\n".join(linhas_md)


def main():
    parser = argparse.ArgumentParser(
        description="Inspeciona planilhas (.xlsx) e amostras (.csv) de Governança de Dados (Gate 0.5 PRODEB / MongoDB / AWS SageMaker)."
    )
    parser.add_argument(
        "caminho",
        nargs="?",
        default=str(DEFAULT_ENTREGA_DIR),
        help="Caminho de um arquivo (.xlsx/.csv) ou diretório a inspecionar (padrão: entrega/governanca_dados/)."
    )
    parser.add_argument(
        "--salvar-relatorio",
        dest="salvar_relatorio",
        default=None,
        help="Caminho opcional para salvar o Relatório de Auditoria/Devolução em Markdown (.md)."
    )
    args = parser.parse_args()

    alvo = Path(args.caminho)
    if not alvo.exists():
        print(f"❌ Caminho não encontrado: {alvo}")
        sys.exit(1)

    arquivos = []
    if alvo.is_file():
        arquivos = [alvo]
    else:
        for ext in ("*.xlsx", "*.ods", "*.csv", "*.pdf", "*.docx", "*.odt"):
            arquivos.extend(sorted(alvo.rglob(ext)))
        arquivos = [f for f in arquivos if not f.name.startswith("~$") and f.name.lower() != "readme.md"]

    if not arquivos:
        print(f"ℹ️ Nenhum arquivo .xlsx, .ods, .csv, .pdf, .docx ou .odt encontrado em: {alvo}")
        print("   Deposite a planilha padrão PRODEB (.xlsx/.ods) e/ou amostras (.csv) em 'entrega/governanca_dados/'.")
        sys.exit(0)

    resultados = []
    for arq in arquivos:
        suf = arq.suffix.lower()
        if suf == ".xlsx":
            resultados.append(inspecionar_xlsx(arq))
        elif suf == ".ods":
            resultados.append(inspecionar_ods(arq))
        elif suf == ".csv":
            resultados.append(inspecionar_csv(arq))
        elif suf == ".pdf":
            resultados.append(inspecionar_pdf(arq))
        elif suf in (".docx", ".odt"):
            resultados.append(inspecionar_doc_texto(arq))

    relatorio_md = gerar_markdown_relatorio(resultados)
    print(relatorio_md)

    if args.salvar_relatorio:
        out_path = Path(args.salvar_relatorio)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(relatorio_md, encoding="utf-8")
        print(f"\n💾 Relatório salvo em: {out_path}")
        if exportar_markdown_para_docx is not None:
            try:
                docx_path = out_path.with_suffix(".docx")
                exportar_markdown_para_docx(
                    str(out_path),
                    str(docx_path),
                    titulo="Relatório de Auditoria de Governança & Devolução ao Cliente (Gate 0.5)",
                )
                print(f"📄 Relatório Word (.docx) salvo em: {docx_path}")
            except Exception:
                pass
        if exportar_markdown_para_pdf is not None:
            try:
                pdf_path = out_path.with_suffix(".pdf")
                exportar_markdown_para_pdf(
                    str(out_path),
                    str(pdf_path),
                    titulo="Relatorio de Auditoria de Governanca & Devolucao ao Cliente (Gate 0.5)",
                )
                print(f"📑 Relatório PDF Oficial (.pdf) salvo em: {pdf_path}")
            except Exception:
                pass

    total_bloqueios = sum(len(r["bloqueios_devolucao"]) for r in resultados)
    sys.exit(1 if total_bloqueios > 0 else 0)


if __name__ == "__main__":
    main()
